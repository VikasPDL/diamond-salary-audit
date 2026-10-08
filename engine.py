"""Salary audit engine: Employee Master + 4 monthly files -> reconciliation workbook."""
import calendar
import difflib
import io
import re
from collections import defaultdict
from datetime import datetime

import openpyxl
import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import attendance_pdf

NO_PROD = "No production"

# Departments that are not expected to show production in DAP
STAFF_DEPTS = {"ADMIN", "H.O.D", "PEON", "SECURITY", "CANTEEN", "CLEANER", "MAINTANCE",
               "ACCOUNTS", "PANDIT JI", "BOILING", "ELECTRICAN", "IT"}


# ---------------------------------------------------------------- helpers
def norm(name):
    """Upper-case, keep letters/digits, single spaces."""
    if name is None:
        return ""
    s = re.sub(r"[^A-Z0-9 ]", " ", str(name).upper())
    return re.sub(r"\s+", " ", s).strip()


def compact(name):
    return norm(name).replace(" ", "")


def num(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def name_score(a, b):
    """0..1 similarity; 'SUNIL MALEKAR' vs 'SUNIL BHAGOJI MALEKAR' scores high (all words contained)."""
    ka, kb = compact(a), compact(b)
    if not ka or not kb:
        return 0.0
    if ka == kb:
        return 1.0
    score = difflib.SequenceMatcher(None, ka, kb).ratio()
    ta = {w for w in norm(a).split() if len(w) > 1}  # ignore initials
    tb = {w for w in norm(b).split() if len(w) > 1}
    small, big = (ta, tb) if len(ta) <= len(tb) else (tb, ta)
    if len(small) >= 2 and small <= big:
        score = max(score, 0.93)
    return score


def best_fuzzy(name, candidates, cutoff=0.85):
    """Return (candidate, score) of best fuzzy match or (None, best score)."""
    best, score = None, 0.0
    for c in candidates:
        r = name_score(name, c)
        if r > score:
            best, score = c, r
    return (best, score) if score >= cutoff else (None, score)


# ---------------------------------------------------------------- readers
def read_master(file):
    wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
    ws = wb["Employee Master"] if "Employee Master" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(hdr)}

    def g(r, col):
        i = idx.get(col)
        return r[i] if i is not None and i < len(r) else None

    out = []
    for r in rows[1:]:
        if not g(r, "Employee Name"):
            continue
        code = g(r, "Emp Code")
        out.append({
            "code": int(code) if isinstance(code, (int, float)) else None,
            "name": str(g(r, "Employee Name")).strip(),
            "designation": g(r, "Designation"),
            "department": g(r, "Department"),
            "sheet": str(g(r, "Sheet Name") or "").strip(),
            "salary_master": g(r, "Salary Master"),
            "bank_name": g(r, "Name in Bank Paid file"),
            "prod_name": g(r, "Name in Employee Production file"),
        })
    return out


def read_salary_sheet(file):
    """All sheets of the monthly salary workbook -> list of employee rows."""
    wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hdr_i = None
        for i, r in enumerate(rows[:15]):
            vals = [str(v).strip().upper() if v is not None else "" for v in r]
            if "NAME" in vals and "PAID" in vals:
                hdr_i, hdr = i, vals
                break
        if hdr_i is None:
            continue

        def col(label, last=False, contains=False):
            hits = [j for j, h in enumerate(hdr) if (label in h if contains else h == label)]
            return (hits[-1] if last else hits[0]) if hits else None

        c_name = col("NAME")
        c_code = col("CODE NO") if col("CODE NO") is not None else col("PUNCH NO")
        if c_code is None:
            c_code = c_name - 1
        c = {"salary": col("SALARY"), "days": col("DAYS"), "abs": col("ABS"),
             "sun": col("SUNDAY", contains=True), "paid": col("PAID"),
             "net": col("NET", last=True, contains=True)}
        for r in rows[hdr_i + 1:]:
            name = r[c_name] if c_name < len(r) else None
            sal = r[c["salary"]] if c["salary"] is not None and c["salary"] < len(r) else None
            code = r[c_code] if c_code < len(r) else None
            srno = r[0] if r else None
            if not isinstance(name, str) or not name.strip():
                continue
            if not isinstance(sal, (int, float)) and not isinstance(srno, (int, float)) and not isinstance(code, (int, float)):
                continue  # section titles / blank lines
            rec = {"sheet": ws.title.strip(), "code": int(code) if isinstance(code, (int, float)) else None,
                   "name": name.strip()}
            for k, j in c.items():
                rec[k] = num(r[j]) if j is not None and j < len(r) else 0.0
            out.append(rec)
    return out


def read_bank(file):
    """Bank transfer file -> {normalised name: {name, amount, accounts}} using the first sheet with a header."""
    wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
    totals = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hdr_i = None
        for i, r in enumerate(rows[:10]):
            vals = [str(v).strip().upper() if v is not None else "" for v in r]
            if "NAME" in vals and "AMOUNT" in vals:
                hdr_i, hdr = i, vals
                break
        if hdr_i is None:
            continue  # numbered split sheets have no header -> skip (they duplicate the main sheet)
        cn, ca = hdr.index("NAME"), hdr.index("AMOUNT")
        cacc = next((j for j, h in enumerate(hdr) if "ACCOUNT" in h), None)
        for r in rows[hdr_i + 1:]:
            if ca >= len(r) or not isinstance(r[cn], str) or not isinstance(r[ca], (int, float)):
                continue
            k = norm(r[cn])
            rec = totals.setdefault(k, {"name": r[cn].strip(), "amount": 0, "accounts": []})
            rec["amount"] += r[ca]
            acc = str(r[cacc]).strip() if cacc is not None and r[cacc] is not None else ""
            if not re.fullmatch(r"\d{6,}", acc):
                acc = ""  # placeholders such as 'NEFT' are not account numbers
            if acc and acc not in rec["accounts"]:
                rec["accounts"].append(acc)
        break  # only the first sheet with a header
    return totals


def read_production(file):
    """Worker production CSV -> {normalised emp_name: dict(name, labour, pcs, cts, rows)}."""
    df = pd.read_csv(file, usecols=["emp_name", "iss_pcs", "iss_cts", "labour"])
    df["key"] = df["emp_name"].map(norm)
    g = df.groupby("key").agg(name=("emp_name", "first"), labour=("labour", "sum"),
                              pcs=("iss_pcs", "sum"), cts=("iss_cts", "sum"), rows=("labour", "size"))
    return {k: dict(v, name=str(v["name"]).strip()) for k, v in g.to_dict("index").items()}


# ---------------------------------------------------------------- matching
def build(master, salary_rows, bank, prod, pdf_emps, month_days, period, tolerance=500):
    mon, year = period
    month_no = datetime.strptime(mon, "%b").month
    sundays = [d for d in range(1, month_days + 1) if calendar.weekday(year, month_no, d) == 6]
    holidays = attendance_pdf.detect_holidays(pdf_emps)
    attendance_pdf.summarise(pdf_emps, month_days, sundays, holidays)

    # --- attendance PDF -> master
    # 1) same Emp Code AND name agrees  2) exact name  3) fuzzy name.  PDF codes are not always
    #    the salary codes, so a code hit with a different name is not trusted.
    pdf_by_code = defaultdict(list)
    for e in pdf_emps:
        pdf_by_code[e["code"]].append(e)
    used_pdf = set()

    def take_pdf(m, e, how):
        m["pdf"], m["pdf_how"] = e, how
        used_pdf.add(id(e))

    for m in master:
        m["pdf"], m["pdf_how"] = None, ""
        for e in pdf_by_code.get(m["code"], []) if m["code"] is not None else []:
            shared = {w for w in norm(m["name"]).split() if len(w) > 3} & set(norm(e["name"]).split())
            if id(e) not in used_pdf and (name_score(m["name"], e["name"]) >= 0.6 or shared):
                take_pdf(m, e, "Emp Code + Name")
                break
    pdf_by_name = defaultdict(list)
    for e in pdf_emps:
        pdf_by_name[compact(e["name"])].append(e)
    for m in master:
        if m["pdf"] is None:
            free = [e for e in pdf_by_name.get(compact(m["name"]), []) if id(e) not in used_pdf]
            if free:
                take_pdf(m, free[0], "Name (exact)" + (", code differs" if free[0]["code"] != m["code"] else ""))
    for m in master:
        if m["pdf"] is None:
            free = [e for e in pdf_emps if id(e) not in used_pdf]
            pick, sc = best_fuzzy(m["name"], [e["name"] for e in free], 0.88)
            if pick:
                take_pdf(m, next(e for e in free if e["name"] == pick), f"Name ({sc:.0%})")

    # --- salary sheet -> master (sheet+code, sheet+name, then name anywhere)
    by_sheet_code = {(s["sheet"].upper(), s["code"]): s for s in salary_rows if s["code"] is not None}
    # when a name is listed twice, prefer the row that has paid days
    ordered = sorted(salary_rows, key=lambda s: (s["paid"] > 0, s["code"] is not None))
    by_sheet_name = {(s["sheet"].upper(), compact(s["name"])): s for s in ordered}
    by_name = {compact(s["name"]): s for s in ordered}
    used_sal = set()
    for m in master:
        sk = m["sheet"].upper()
        s, how = by_sheet_code.get((sk, m["code"])), "Sheet + Code"
        if s is None or id(s) in used_sal:
            s, how = by_sheet_name.get((sk, compact(m["name"]))), "Sheet + Name"
        if s is None or id(s) in used_sal:
            s, how = by_name.get(compact(m["name"])), "Name"
        if s is not None and id(s) in used_sal:
            s = None
        m["sal"], m["sal_how"] = s, (how if s else "")
        if s is not None:
            used_sal.add(id(s))

    # --- production + bank: master mapping column, else exact master name
    used_prod, used_bank = set(), set()
    for m in master:
        m["prod"], m["prod_how"] = None, ""
        for label, nm in (("Master mapping", m["prod_name"]), ("Employee name", m["name"])):
            if nm and norm(nm) in prod:
                m["prod"], m["prod_how"] = prod[norm(nm)], label
                used_prod.add(norm(nm))
                break
        m["bank"], m["bank_how"] = None, ""
        for label, nm in (("Master mapping", m["bank_name"]), ("Employee name", m["name"])):
            if nm and norm(nm) in bank and norm(nm) not in used_bank:
                m["bank"], m["bank_how"] = bank[norm(nm)], label
                used_bank.add(norm(nm))
                break

    # --- employees in salary sheet but missing from master -> add as extra rows
    extra = []
    matched_names = {compact(m["name"]) for m in master if m["sal"]}
    for s in salary_rows:
        if id(s) in used_sal:
            continue
        dup = compact(s["name"]) in matched_names
        e = {"code": s["code"], "name": s["name"], "designation": None, "department": None,
             "sheet": s["sheet"], "salary_master": s["salary"], "bank_name": None, "prod_name": None,
             "sal": s, "sal_how": "Not in master", "pdf": None, "prod": None, "bank": None,
             "prod_how": "", "bank_how": "", "new": True, "dup": dup}
        if dup:  # second line for the same person in the salary sheet: show it, don't re-link other files
            extra.append(e)
            continue
        free = [x for x in pdf_emps if id(x) not in used_pdf]
        pick, sc = best_fuzzy(s["name"], [x["name"] for x in free], 0.88)
        if pick:
            take_pdf(e, next(x for x in free if x["name"] == pick), f"Name ({sc:.0%})")
        if norm(s["name"]) in prod:
            e["prod"], e["prod_how"] = prod[norm(s["name"])], "Employee name"
            used_prod.add(norm(s["name"]))
        if norm(s["name"]) in bank and norm(s["name"]) not in used_bank:
            e["bank"], e["bank_how"] = bank[norm(s["name"])], "Employee name"
            used_bank.add(norm(s["name"]))
        extra.append(e)

    people = master + extra

    # --- salary paid into a bank account under another name
    # e.g. 'MAMTA RAJAN KASHALKAR' (unlinked bank name) shares 2+ words with 'RAJAN BHALCHANDRA KASHALKAR'
    for p in people:
        p["other_bank"], p["same_acc"] = [], []
    for k, b in bank.items():
        if k in used_bank:
            continue
        words = {w for w in norm(b["name"]).split() if len(w) > 2}
        scored = [(len(words & {w for w in norm(p["name"]).split() if len(w) > 2}), i)
                  for i, p in enumerate(people) if not p.get("dup")]
        best = max((sc for sc, _ in scored), default=0)
        hits = [i for sc, i in scored if sc == best]
        if best >= 2 and len(hits) == 1:
            people[hits[0]]["other_bank"].append(b)
            used_bank.add(k)
    # one account number used for different names
    acc_names = defaultdict(set)
    for b in bank.values():
        for a in b["accounts"]:
            acc_names[a].add(b["name"])
    for p in people:
        for b in ([p["bank"]] if p["bank"] else []) + p["other_bank"]:
            for a in b["accounts"]:
                others = sorted(acc_names[a] - {b["name"]})
                if others:
                    p["same_acc"].append((a, others))

    # --- figures + remarks
    for p in people:
        p["dap"] = round(p["prod"]["labour"], 2) if p["prod"] else NO_PROD
        p["pcs"] = int(p["prod"]["pcs"]) if p["prod"] else 0
        p["cts"] = round(p["prod"]["cts"], 2) if p["prod"] else 0
        p["bank_amt"] = (p["bank"]["amount"] if p["bank"] else 0) + sum(b["amount"] for b in p["other_bank"])
        p["att_pdf"] = p["pdf"]["paid"] if p["pdf"] else 0
        p["att_sal"] = p["sal"]["paid"] if p["sal"] else 0
        p["net_sal"] = round(p["sal"]["net"], 2) if p["sal"] else 0
        p["att_diff"] = p["att_pdf"] - p["att_sal"]
        p["eff"] = round(p["dap"] - p["bank_amt"], 2) if p["dap"] != NO_PROD else 0
        p["remark_extra"] = remark_extra(p)
        p["remarks"] = base_remark(p["dap"], p["att_diff"]) + p["remark_extra"]

    unmatched = {
        "pdf": [e for e in pdf_emps if id(e) not in used_pdf],
        "prod": [dict(v, key=k) for k, v in prod.items() if k not in used_prod],
        "bank": [{"name": v["name"], "amount": v["amount"], "account": ", ".join(v["accounts"])}
                 for k, v in bank.items() if k not in used_bank],
    }
    # suggestions for unmatched production / bank names
    all_names = [p["name"] for p in people]
    for lst in (unmatched["prod"], unmatched["bank"]):
        for u in lst:
            u["suggest"], u["score"] = best_fuzzy(u["name"], all_names, 0.8)
    meta = {"period": f"{mon}-{year}", "month_days": month_days, "sundays": sundays,
            "holidays": holidays, "tolerance": tolerance}
    return people, unmatched, meta


REMARKS = ["High Paid", "OK", "No Production"]


def base_remark(dap, att_diff):
    """No production in DAP -> No Production; Diff - Attendance <= 0 -> High Paid; else OK."""
    if dap == NO_PROD:
        return "No Production"
    return "High Paid" if att_diff <= 0 else "OK"


def remark_extra(p):
    """Bank details appended to the remark: salary paid to another name / shared account number."""
    parts = []
    for b in p["other_bank"]:
        parts.append(f"Salary also paid to other A/c: {b['name']} A/c {', '.join(b['accounts']) or '-'} Rs {b['amount']:,.0f}")
    for acc, others in p["same_acc"]:
        parts.append(f"Same A/c {acc} also used for {', '.join(others)}")
    return "".join("; " + x for x in parts)


# ---------------------------------------------------------------- workbook
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header(ws, row, headers, widths=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name=FONT, bold=True)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 45
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=3)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"


def _body_style(ws, first, last, ncols, numfmt=None):
    for r in ws.iter_rows(min_row=first, max_row=last, max_col=ncols):
        for c in r:
            c.font = Font(name=FONT)
            c.border = BORDER
            if numfmt and c.column in numfmt:
                c.number_format = numfmt[c.column]


def write_workbook(people, unmatched, meta):
    wb = openpyxl.Workbook()
    money = "#,##0;(#,##0);-"

    # ---------- Sheet 1: main view
    ws = wb.active
    ws.title = "Diamond Report"
    hdr = ["Emp Code", "Employee Name", "Designation", "Salary Master", "Earn Salary - DAP", "Pcs Done",
           "Cts Done", "Diff - Attdance", "Effiency - Bank trans - Salry DAP", "Remarks"]
    _header(ws, 1, hdr, [9, 34, 22, 12, 15, 10, 11, 11, 17, 55])
    # Supporting figures referenced by the formulas live on the 'Calculation' sheet (same row order)
    for i, p in enumerate(people, 2):
        ws.cell(row=i, column=1, value=p["code"])
        ws.cell(row=i, column=2, value=p["name"])
        ws.cell(row=i, column=3, value=p["designation"])
        ws.cell(row=i, column=4, value=p["salary_master"])
        ws.cell(row=i, column=5, value=p["dap"])
        ws.cell(row=i, column=6, value=p["pcs"])
        ws.cell(row=i, column=7, value=p["cts"])
        ws.cell(row=i, column=8, value=f"=Calculation!J{i}-Calculation!K{i}")
        ws.cell(row=i, column=9, value=f"=IF(ISNUMBER(E{i}),E{i}-Calculation!H{i},0)")
        extra = p["remark_extra"].replace('"', "'")
        ws.cell(row=i, column=10, value=f'=IF(ISNUMBER(E{i}),IF(H{i}<=0,"High Paid","OK"),"No Production")'
                                        + (f'&"{extra}"' if extra else ""))
    last = len(people) + 1
    _body_style(ws, 2, last, len(hdr), {4: money, 5: money, 6: "#,##0", 7: "#,##0.00", 8: "0.0;-0.0;-", 9: money})
    for i, p in enumerate(people, 2):
        if p.get("new"):
            for j in range(1, len(hdr) + 1):
                ws.cell(row=i, column=j).fill = WARN_FILL
    ws.cell(row=1, column=8).comment = Comment("Paid days as per attendance PDF minus paid days as per salary sheet", "Audit")
    ws.cell(row=1, column=9).comment = Comment("Earn Salary - DAP minus Bank Transfer amount (0 when no production)", "Audit")
    ws.cell(row=1, column=10).comment = Comment("No Production = no DAP production; High Paid = Diff - Attdance <= 0; "
                                                "otherwise OK. Bank details added when salary also went to another name's account.", "Audit")

    # ---------- Calculation backup
    wc = wb.create_sheet("Calculation")
    chdr = ["Emp Code", "Employee Name", "Department", "Salary Sheet Tab", "Salary Master",
            "Earn Salary - DAP", "Net Salary (Salary Sheet)", "Bank Transfer (incl. other-name A/c)", "Diff - Salary Sheet vs Bank",
            "Paid Days (Attendance PDF)", "Paid Days (Salary Sheet)", "OT Hrs (PDF)", "Pcs Done", "Cts Done"]
    _header(wc, 1, chdr, [9, 34, 24, 16, 12, 14, 14, 14, 14, 12, 12, 10, 10, 11])
    for i, p in enumerate(people, 2):
        vals = [p["code"], p["name"], p["department"] or p["sheet"] or "Unassigned", p["sheet"], p["salary_master"], p["dap"], p["net_sal"],
                p["bank_amt"], f"=G{i}-H{i}", p["att_pdf"], p["att_sal"],
                p["pdf"]["ot_hrs"] if p["pdf"] else 0, p["pcs"], p["cts"]]
        for j, v in enumerate(vals, 1):
            wc.cell(row=i, column=j, value=v)
    _body_style(wc, 2, last, len(chdr), {5: money, 6: money, 7: money, 8: money, 9: money, 13: "#,##0", 14: "#,##0.00"})

    # ---------- Sheet 2: Attendance P/A
    wa = wb.create_sheet("Attendance PA")
    days = list(range(1, meta["month_days"] + 1))
    ahdr = (["Emp Code", "Employee Name", "Department (PDF)"] + [str(d) for d in days] +
            ["Present", "Absent (A)", "Half Day", "ABS Days", "Sunday & Paid Holiday", "Paid Days (PDF)",
             "Paid Days (Salary Sheet)", "Diff", "OT Hrs"])
    wa.cell(row=1, column=1, value=f"Attendance {meta['period']} - from attendance PDF. "
            f"Sundays: {', '.join(map(str, meta['sundays']))}; Holidays detected: {', '.join(map(str, meta['holidays'])) or 'none'}. "
            "ABS = A + 0.5 x HD; Sunday & Paid Holiday = days worked on a Sunday/holiday (or WP/HWP/PH/HDH); "
            "Paid Days = Month days - ABS + Sunday & Paid Holiday.").font = Font(name=FONT, italic=True)
    _header(wa, 2, ahdr, [9, 32, 20] + [5] * len(days) + [9, 9, 9, 9, 11, 10, 10, 8, 8])
    pdf_people = [p for p in people if p["pdf"]] + [
        {"code": e["code"], "name": e["name"], "pdf": e, "sal": None, "att_sal": 0, "unmatched": True}
        for e in unmatched["pdf"]]
    n0 = 3 + len(days)
    for i, p in enumerate(pdf_people, 3):
        e = p["pdf"]
        wa.cell(row=i, column=1, value=p["code"])
        wa.cell(row=i, column=2, value=p["name"] + ("  (not in master)" if p.get("unmatched") else ""))
        wa.cell(row=i, column=3, value=e["department"])
        for j, d in enumerate(days, 4):
            code = e["days"].get(d, "")
            c = wa.cell(row=i, column=j, value=code)
            c.alignment = Alignment(horizontal="center")
            if code == "A":
                c.fill = WARN_FILL
            elif code in ("WD", "H"):
                c.fill = HDR_FILL
        cnt = e["counts"]
        wa.cell(row=i, column=n0 + 1, value=e["present"])
        wa.cell(row=i, column=n0 + 2, value=cnt["A"])
        wa.cell(row=i, column=n0 + 3, value=cnt["HD"])
        wa.cell(row=i, column=n0 + 4, value=e["abs"])
        wa.cell(row=i, column=n0 + 5, value=e["sun_ph"])
        L = get_column_letter
        wa.cell(row=i, column=n0 + 6, value=f"={meta['month_days']}-{L(n0 + 4)}{i}+{L(n0 + 5)}{i}")
        wa.cell(row=i, column=n0 + 7, value=p["att_sal"] if p.get("sal") else None)
        wa.cell(row=i, column=n0 + 8, value=f'=IF({L(n0 + 7)}{i}="","",{L(n0 + 6)}{i}-{L(n0 + 7)}{i})')
        wa.cell(row=i, column=n0 + 9, value=e["ot_hrs"])
    _body_style(wa, 3, len(pdf_people) + 2, len(ahdr))
    wa.freeze_panes = "D3"

    # ---------- Sheet 3: Name master matching
    wm = wb.create_sheet("Name Matching")
    mhdr = ["Emp Code", "Master Employee Name", "Department", "Salary Sheet Tab",
            "Name in Salary Sheet", "Salary Sheet Match", "Name in Production (DAP)", "Production Match",
            "Name in Bank File", "Bank Match", "Name in Attendance PDF", "Attendance Match", "Status"]
    _header(wm, 1, mhdr, [9, 32, 22, 16, 32, 14, 32, 15, 32, 15, 32, 14, 30])
    for i, p in enumerate(people, 2):
        missing = [lbl for lbl, k in (("Salary", "sal"), ("Bank", "bank"), ("Attendance", "pdf"))
                   if not p[k] and not (k == "bank" and p["other_bank"])]
        status = "Duplicate line in Salary Sheet" if p.get("dup") else "Not in Employee Master" if p.get("new") else ("All matched" if not missing else "Missing: " + ", ".join(missing))
        vals = [p["code"], p["name"], p["department"] or p["sheet"] or "Unassigned", p["sheet"],
                p["sal"]["name"] if p["sal"] else "", p["sal_how"] or "Not found",
                p["prod"]["name"] if p["prod"] else "", p["prod_how"] or "No production",
                "; ".join(([p["bank"]["name"]] if p["bank"] else []) + [b["name"] + " (other name)" for b in p["other_bank"]]),
                p["bank_how"] or ("Other name" if p["other_bank"] else "Not found"),
                p["pdf"]["name"] if p["pdf"] else "", p.get("pdf_how") or "Not found", status]
        for j, v in enumerate(vals, 1):
            wm.cell(row=i, column=j, value=v)
    _body_style(wm, 2, last, len(mhdr))
    for i, p in enumerate(people, 2):
        wm.cell(row=i, column=13).fill = OK_FILL if wm.cell(row=i, column=13).value == "All matched" else WARN_FILL

    # ---------- Unmatched source names
    wu = wb.create_sheet("Unmatched Names")
    r = 1
    for title, cols, rows in (
        ("Production (DAP) names not linked to any employee", ["Name in Production", "Labour (Rs)", "Pcs", "Cts", "Suggested Master Name", "Similarity"],
         [[u["name"], round(u["labour"], 2), u["pcs"], round(u["cts"], 2), u["suggest"] or "", round(u["score"], 2)]
          for u in sorted(unmatched["prod"], key=lambda u: -u["labour"])]),
        ("Bank transfer names not linked to any employee", ["Name in Bank File", "Amount", "Account No", "Suggested Master Name", "Similarity"],
         [[u["name"], u["amount"], u["account"], u["suggest"] or "", round(u["score"], 2)]
          for u in sorted(unmatched["bank"], key=lambda u: -u["amount"])]),
        ("Attendance PDF employees not linked to any employee", ["Emp Code", "Name in PDF", "Department", "Paid Days"],
         [[e["code"], e["name"], e["department"], e["paid"]] for e in unmatched["pdf"]]),
    ):
        wu.cell(row=r, column=1, value=f"{title} ({len(rows)})").font = Font(name=FONT, bold=True, size=12)
        r += 1
        for j, h in enumerate(cols, 1):
            c = wu.cell(row=r, column=j, value=h)
            c.font, c.fill, c.border = Font(name=FONT, bold=True), HDR_FILL, BORDER
        for row in rows:
            r += 1
            for j, v in enumerate(row, 1):
                c = wu.cell(row=r, column=j, value=v)
                c.font, c.border = Font(name=FONT), BORDER
        r += 3
    for j, w in enumerate([34, 14, 30, 30, 34, 11], 1):
        wu.column_dimensions[get_column_letter(j)].width = w

    # ---------- Remarks pivot
    wr = wb.create_sheet("Remarks", 1)
    bold = Font(name=FONT, bold=True)
    wr.cell(row=1, column=1, value=f"Remarks summary - {meta['period']}").font = Font(name=FONT, bold=True, size=13)
    wr.cell(row=2, column=1, value="No Production = no DAP production; High Paid = Diff - Attdance <= 0; OK = rest.").font = Font(name=FONT, italic=True)
    R = "'Diamond Report'!$J:$J"
    _header(wr, 4, ["Remarks", "Count", "Earn Salary - DAP", "Effiency - Bank trans - Salry DAP"], [36, 14, 16, 16, 16, 12])
    wr.freeze_panes = None
    wr.auto_filter.ref = None
    for i, k in enumerate(REMARKS, 5):
        wr.cell(row=i, column=1, value=k)
        wr.cell(row=i, column=2, value=f'=COUNTIF({R},A{i}&"*")')
        wr.cell(row=i, column=3, value=f"=SUMIF({R},A{i}&\"*\",'Diamond Report'!$E:$E)")
        wr.cell(row=i, column=4, value=f"=SUMIF({R},A{i}&\"*\",'Diamond Report'!$I:$I)")
    t = 5 + len(REMARKS)
    wr.cell(row=t, column=1, value="Grand Total")
    for col in (2, 3, 4):
        L = get_column_letter(col)
        wr.cell(row=t, column=col, value=f"=SUM({L}5:{L}{t - 1})")
    wr.cell(row=t + 1, column=1, value="Employees in Diamond Report")
    wr.cell(row=t + 1, column=2, value="=COUNTA('Diamond Report'!$B:$B)-1")
    wr.cell(row=t + 2, column=1, value="Salary also paid to other-name A/c")
    wr.cell(row=t + 2, column=2, value=f'=COUNTIF({R},"*other A/c*")')
    _body_style(wr, 5, t + 2, 4, {3: money, 4: money})
    for row in (t, t + 1, t + 2):
        for col in (1, 2, 3, 4):
            wr.cell(row=row, column=col).font = bold

    # pivot: department x remark
    top = t + 5
    groups = sorted({str(p["department"] or p["sheet"] or "Unassigned") for p in people})
    phdr = ["Department"] + REMARKS + ["Total"]
    for j, h in enumerate(phdr, 1):
        c = wr.cell(row=top, column=j, value=h)
        c.font, c.fill, c.border = bold, HDR_FILL, BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, g in enumerate(groups, top + 1):
        wr.cell(row=i, column=1, value=g)
        for j, k in enumerate(REMARKS, 2):
            wr.cell(row=i, column=j, value=f'=COUNTIFS(Calculation!$C:$C,$A{i},{R},"{k}*")')
        wr.cell(row=i, column=len(phdr), value=f"=SUM(B{i}:{get_column_letter(len(phdr) - 1)}{i})")
    gt = top + 1 + len(groups)
    wr.cell(row=gt, column=1, value="Grand Total")
    for j in range(2, len(phdr) + 1):
        L = get_column_letter(j)
        wr.cell(row=gt, column=j, value=f"=SUM({L}{top + 1}:{L}{gt - 1})")
    _body_style(wr, top + 1, gt, len(phdr))
    for j in range(1, len(phdr) + 1):
        wr.cell(row=gt, column=j).font = bold

    # salary paid to other-name accounts / shared account numbers
    ob = [(p, b) for p in people for b in p["other_bank"]]
    sa = [(p, a, o) for p in people for a, o in p["same_acc"]]
    r0 = gt + 3
    wr.cell(row=r0, column=1, value=f"Salary paid to other-name bank account ({len(ob)})").font = Font(name=FONT, bold=True, size=12)
    for j, h in enumerate(["Employee Name", "Emp Code", "Paid to (name in bank file)", "Account No", "Amount"], 1):
        c = wr.cell(row=r0 + 1, column=j, value=h)
        c.font, c.fill, c.border = bold, HDR_FILL, BORDER
    for i, (p, b) in enumerate(ob, r0 + 2):
        for j, v in enumerate([p["name"], p["code"], b["name"], ", ".join(b["accounts"]), b["amount"]], 1):
            c = wr.cell(row=i, column=j, value=v)
            c.font, c.border = Font(name=FONT), BORDER
        wr.cell(row=i, column=5).number_format = money
    r1 = r0 + 4 + len(ob)
    wr.cell(row=r1, column=1, value=f"Same account number used for different names ({len(sa)})").font = Font(name=FONT, bold=True, size=12)
    for j, h in enumerate(["Employee Name", "Emp Code", "Account No", "Also used for"], 1):
        c = wr.cell(row=r1 + 1, column=j, value=h)
        c.font, c.fill, c.border = bold, HDR_FILL, BORDER
    for i, (p, a, o) in enumerate(sa, r1 + 2):
        for j, v in enumerate([p["name"], p["code"], a, ", ".join(o)], 1):
            c = wr.cell(row=i, column=j, value=v)
            c.font, c.border = Font(name=FONT), BORDER
    for j, w in enumerate([36, 14, 34, 22, 16, 12], 1):
        wr.column_dimensions[get_column_letter(j)].width = w

    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run(master_file, salary_file, bank_file, prod_file, pdf_file, tolerance=500):
    master = read_master(master_file)
    salary_rows = read_salary_sheet(salary_file)
    bank = read_bank(bank_file)
    prod = read_production(prod_file)
    pdf_emps, month_days, period = attendance_pdf.parse_pdf(pdf_file)
    people, unmatched, meta = build(master, salary_rows, bank, prod, pdf_emps, month_days, period, tolerance)
    return people, unmatched, meta, write_workbook(people, unmatched, meta)
